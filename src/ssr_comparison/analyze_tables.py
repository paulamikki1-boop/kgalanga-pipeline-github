import os
OUT_DIR = os.environ.get("KG_OUTPUT_DIR", "output")
IN_DIR = os.environ.get("KG_INPUT_DIR", "data")
os.makedirs(OUT_DIR, exist_ok=True)
import openpyxl, json
from collections import defaultdict

path = IN_DIR + "/ssr of kb and ks.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)

# ---- Table 1 source: regions ----
ws = wb["regions"]
regions_raw = []
for r in range(1, ws.max_row + 1):
    row = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
    if any(v is not None for v in row):
        regions_raw.append(row)

region_data = [
    {"region": "Large single copy (LSC)", "size": 88429, "gc": 34},
    {"region": "Inverted repeat B (IRB)", "size": 29797, "gc": 41},
    {"region": "Small single copy (SSC)", "size": 15776, "gc": 30},
    {"region": "Inverted repeat A (IRA)", "size": 29797, "gc": 41},
]
total_size = sum(r["size"] for r in region_data)
weighted_gc = sum(r["size"] * r["gc"] for r in region_data) / total_size

# ---- CDS totals from ks tab (dedupe exact duplicate rows) ----
ws = wb["ks"]
seen = set()
cds_rows = []
gene_rows = []
trna_rows = []
rrna_rows = []
for r in range(2, ws.max_row + 1):
    region = ws.cell(row=r, column=1).value
    feat = ws.cell(row=r, column=2).value
    start = ws.cell(row=r, column=4).value
    end = ws.cell(row=r, column=5).value
    size = ws.cell(row=r, column=6).value
    typ = ws.cell(row=r, column=7).value
    key = (feat, typ, start, end, size)
    if key in seen:
        continue
    seen.add(key)
    if typ == "CDS":
        cds_rows.append({"gene": feat, "region": region, "size": size, "start": start, "end": end})
    if typ == "rRNA":
        rrna_rows.append({"gene": feat, "region": region, "size": size})

unique_cds_genes = sorted(set(r["gene"] for r in cds_rows))
# total protein-coding bp: sum CDS sizes but for genes with a genomic-span CDS row that
# duplicates an already-counted exon/intron total (e.g. petD 1229 = 8+746+475 already
# counted via exon/intron rows), keep the LARGEST CDS size per (gene, region) as the
# representative length for that locus copy, to avoid double counting.
by_locus = defaultdict(list)
for r in cds_rows:
    if r["gene"] == "rps12" or (r["end"] is not None and r["start"] is not None and r["end"] < r["start"]):
        continue  # corrupted join()-location parse (trans-spliced rps12); handled separately
    by_locus[(r["gene"], r["region"])].append(r["size"])
protein_coding_bp = sum(max(v) for v in by_locus.values())
n_protein_coding_loci = len(unique_cds_genes)  # includes rps12
n_protein_coding_copies = len(by_locus) + 2  # + rps12's 2 IR copies, excluded from bp sum above

rrna_unique = sorted(set(r["gene"] for r in rrna_rows))
rrna_bp = sum(max(s for g2, s in [(x["gene"], x["size"]) for x in rrna_rows] if g2 == g) for g in set(r["gene"] for r in rrna_rows))
# simpler: sum by locus copy (region,gene) taking size once per copy
rrna_by_locus = defaultdict(list)
for r in rrna_rows:
    rrna_by_locus[(r["gene"], r["region"])].append(r["size"])
rrna_bp = sum(max(v) for v in rrna_by_locus.values())

print("Total region size check:", total_size)
print("Weighted genome GC%:", round(weighted_gc, 1))
print("Unique protein-coding gene names:", n_protein_coding_loci)
print("Protein-coding physical locus copies (incl. IR dup):", n_protein_coding_copies)
print("Protein-coding total bp (deduped):", protein_coding_bp)
print("rRNA unique genes:", rrna_unique)
print("rRNA physical copies:", len(rrna_by_locus))
print("rRNA total bp:", rrna_bp)

# ---- tRNA table from genes tab ----
ws = wb["genes"]
trna = []
for r in range(3, 33):
    name = ws.cell(row=r, column=5).value
    anti = ws.cell(row=r, column=6).value
    cnt = ws.cell(row=r, column=7).value
    strand = ws.cell(row=r, column=10).value
    note = ws.cell(row=r, column=11).value
    if anti and cnt is not None:
        trna.append({"aa": name, "gene": f"trn{name[0] if name else ''}", "anticodon": anti,
                     "count": cnt, "strand": strand, "note": note})
trna_total = sum(t["count"] for t in trna)
print("tRNA rows:", len(trna), "sum of counts:", trna_total)

# ---- protein-coding categories from genes tab (cols A,B) ----
ws = wb["genes"]
categories = []
for r in range(3, 15):
    cat = ws.cell(row=r, column=1).value
    genes = ws.cell(row=r, column=2).value
    if cat:
        categories.append({"category": cat, "genes": genes})
for c in categories:
    print(c)

# save everything needed for docx build
out = {
    "region_data": region_data,
    "total_size": total_size,
    "weighted_gc": round(weighted_gc, 2),
    "protein_coding_bp": protein_coding_bp,
    "n_protein_coding_loci": n_protein_coding_loci,
    "n_protein_coding_copies": n_protein_coding_copies,
    "rrna_bp": rrna_bp,
    "rrna_unique": rrna_unique,
    "rrna_copies": len(rrna_by_locus),
    "trna": trna,
    "trna_total": trna_total,
    "categories": categories,
    "cds_rows": cds_rows,
}
with open(OUT_DIR + "/tables_data.json", "w") as f:
    json.dump(out, f, indent=2, default=str)
print("saved tables_data.json")

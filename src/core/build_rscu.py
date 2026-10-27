import os
OUT_DIR = os.environ.get("KG_OUTPUT_DIR", "output")
IN_DIR = os.environ.get("KG_INPUT_DIR", "data")
os.makedirs(OUT_DIR, exist_ok=True)
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

data = json.load(open(OUT_DIR + "/rscu_table.json"))
rows = data["rscu_rows"]  # [aa_name, codon(U), count, rscu(precomputed, will replace w/ formula)]

# reorder: 20 standard aa alphabetically (already are, minus Stop), Stop group moved to end
aa_rows = [r for r in rows if r[0] != "Stop"]
stop_rows = [r for r in rows if r[0] == "Stop"]
ordered = aa_rows + stop_rows

wb = Workbook()
ws = wb.active
ws.title = "RSCU"

FONT = "Arial"
header_font = Font(name=FONT, bold=True, color="FFFFFF", size=11)
header_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
normal_font = Font(name=FONT, size=11)
italic_font = Font(name=FONT, size=10, italic=True)
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center")

headers = ["Amino Acid", "Codon", "Codon Count (KB = KS)", "RSCU"]
for col, h in enumerate(headers, start=1):
    c = ws.cell(row=1, column=col, value=h)
    c.font = header_font
    c.fill = header_fill
    c.alignment = center
    c.border = border

start_row = 2
for i, (aa, codon, count, _rscu_precalc) in enumerate(ordered):
    r = start_row + i
    ws.cell(row=r, column=1, value=aa).font = normal_font
    ws.cell(row=r, column=2, value=codon).font = normal_font
    ws.cell(row=r, column=3, value=count).font = normal_font
    # RSCU = n_family * count_i / SUM(counts in family)
    formula = (f"=COUNTIF($A${start_row}:$A${start_row+len(ordered)-1},A{r})*C{r}"
               f"/SUMIF($A${start_row}:$A${start_row+len(ordered)-1},A{r},$C${start_row}:$C${start_row+len(ordered)-1})")
    fc = ws.cell(row=r, column=4, value=formula)
    fc.font = normal_font
    fc.number_format = "0.00"
    for col in range(1,5):
        ws.cell(row=r, column=col).border = border
        ws.cell(row=r, column=col).alignment = center

last_row = start_row + len(ordered) - 1
total_row = last_row + 2
ws.cell(row=total_row, column=1, value="Total codons").font = Font(name=FONT, bold=True, size=11)
tf = ws.cell(row=total_row, column=3, value=f"=SUM(C{start_row}:C{last_row})")
tf.font = Font(name=FONT, bold=True, size=11)

note_row = total_row + 2
notes = [
 "Source: canonical CDS of 78 protein-coding genes (clpP1 excluded), GeSeqJob-galanga_KB_GenBank.gb / KS_GenBank.gb.",
 "KB and KS give one shared table: all 79 annotated protein-coding genes translate identically between the two genomes (see QC report).",
 "Genetic code: table 11 (bacterial/plant plastid). RSCU computed per amino-acid synonymous family; Stop codons shown separately, RSCU computed within the 3-codon stop family.",
 "Where GeSeq produced 2 overlapping CDS calls for a gene (ycf3/pafI, petD, ndhK), the biologically valid reading was used - see KB_KS_Table1_and_QC_report.docx for the resolution.",
]
for i, n in enumerate(notes):
    cell = ws.cell(row=note_row+i, column=1, value=n)
    cell.font = italic_font
    ws.merge_cells(start_row=note_row+i, start_column=1, end_row=note_row+i, end_column=4)

ws.column_dimensions["A"].width = 16
ws.column_dimensions["B"].width = 10
ws.column_dimensions["C"].width = 22
ws.column_dimensions["D"].width = 10
ws.freeze_panes = "A2"

wb.save(OUT_DIR + "/KB_KS_RSCU_codon_usage.xlsx")
print("saved")

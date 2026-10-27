import os
OUT_DIR = os.environ.get("KG_OUTPUT_DIR", "output")
IN_DIR = os.environ.get("KG_INPUT_DIR", "data")
os.makedirs(OUT_DIR, exist_ok=True)
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

rows = json.load(open(OUT_DIR + "/kb_vs_ref_variants_annotated.json"))

wb = Workbook()
ws = wb.active
ws.title = "Variants_vs_MK209001"
FONT = "Arial"
header_font = Font(name=FONT, bold=True, color="FFFFFF", size=11)
header_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
coding_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
normal_font = Font(name=FONT, size=11)
italic_font = Font(name=FONT, size=10, italic=True)
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)

headers = ["Ref. position (1-based, MK209001)","Type","Ref allele","KB=KS allele","Region / gene","In coding sequence?"]
widths = [22, 10, 24, 30, 26, 16]
for col,(h,w) in enumerate(zip(headers,widths), start=1):
    c = ws.cell(row=1, column=col, value=h); c.font=header_font; c.fill=header_fill; c.alignment=center; c.border=border
    ws.column_dimensions[chr(64+col)].width = w

for i, r in enumerate(rows, start=2):
    vals = [r["ref_pos"]+1, r["type"], r["ref_allele"], r["query_allele"], r["region"], "Yes" if r["in_cds"] else "No"]
    for col, v in enumerate(vals, start=1):
        cell = ws.cell(row=i, column=col, value=v)
        cell.font = normal_font
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if r["in_cds"]:
            cell.fill = coding_fill

last = len(rows)+1
notes_row = last+2
notes = [
 "KB and KS are 100% identical to each other genome-wide (see KB_KS_Table1_and_QC_report.docx), so this single variant list applies to both vs the reference.",
 "Positions are on the reference's (MK209001) collinear LSC+IRb+SSC coordinate system (one IR copy only, to avoid double-counting); IRa is the exact duplicate of IRb and carries the same 0 variants.",
 "Whole-genome pairwise alignment: minimap2 (asm5 preset) via the mappy Python binding. 39 variant events total: 22 SNPs, 9 insertions (68 bp), 8 deletions (80 bp) = 148 bp indel.",
 "6 of 39 fall inside an annotated CDS (highlighted): rpoC2, rpl33, rpl16 synonymous; ycf4, rpl14, ycf1 nonsynonymous. See comparison report for codon-level detail and pooled Ka/Ks.",
]
for i,n in enumerate(notes):
    cell = ws.cell(row=notes_row+i, column=1, value=n)
    cell.font = italic_font
    ws.merge_cells(start_row=notes_row+i, start_column=1, end_row=notes_row+i, end_column=6)

ws.freeze_panes = "A2"
wb.save(OUT_DIR + "/KB_KS_vs_MK209001_variants.xlsx")
print("saved")
